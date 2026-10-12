"""
formatos.py — leitura e escrita dos formatos do Artry (Arch Quilly).

O editor trabalha com HTML (contenteditable). Aqui:
  .artry   HTML completo (guarda toda a formatação). .artry antigo (texto puro) também abre.
  .hyp/.md Markdown (Hyper Doc v2: front-matter e anexos preservados) — a conversão HTML->Markdown é
           feita no JavaScript; aqui só Markdown->HTML (python-markdown).
  .rtf     RTF <-> HTML com formatação (negrito, itálico, sublinhado, tachado, cor, realce, tamanho,
           fonte, alinhamento, recuo, espaçamento, acentos/Unicode).
  .docx    DOCX <-> HTML com formatação (python-docx): títulos, listas, alinhamento, runs.
  texto    .txt, .log, código...: texto puro UTF-8.
  preview  as primeiras linhas em texto puro (cartões da tela inicial).
"""
import base64
import html
import os
import re
from html.parser import HTMLParser

EXT_CODIGO = {".cs", ".csx", ".js", ".ts", ".py", ".java", ".cpp", ".c", ".h", ".css", ".sql", ".sh",
              ".ps1", ".json", ".xml", ".html", ".yml", ".yaml", ".htm", ".rs", ".go", ".lua", ".toml",
              ".ini", ".conf", ".gd", ".kt"}
ANEXO_MARCA = "<!-- ===== ANEXOS HYP ===== -->"


def esc(s):
    return html.escape(s or "", quote=False)


# ============================================================ .hyp / markdown
def hyp_partes(texto):
    """-> (front-matter cru ou "", corpo markdown, {id: (mime, base64)})."""
    fm = ""
    corpo, _, rodape = texto.partition(ANEXO_MARCA)
    m = re.match(r"\s*---\r?\n.*?\r?\n---\r?\n", corpo, re.S)
    if m:
        fm, corpo = corpo[:m.end()], corpo[m.end():]
    anexos = {}
    for a in re.finditer(r"```hyp-anexo\s+(\S+)\s+(\S+)\s*\n(.*?)\n```", rodape, re.S):
        anexos[a.group(1).lower()] = (a.group(2), re.sub(r"\s+", "", a.group(3)))
    return fm, corpo, anexos


def markdown_para_html(texto):
    """.hyp/.md -> (html, front-matter). Imagens anexo:id viram data: URIs (voltam a ser anexos ao salvar)."""
    import markdown
    fm, corpo, anexos = hyp_partes(texto)

    def img(m):
        alt, id_ = m.group(1), m.group(2).lower()
        if id_ in anexos:
            mime, b64 = anexos[id_]
            return f"![{alt}](data:{mime};base64,{b64})"
        return f"[{alt}]"
    corpo = re.sub(r"!\[([^\]]*)\]\(anexo:([\w.\-]+)\)", img, corpo, flags=re.I)
    html_ = markdown.markdown(corpo, extensions=["extra", "sane_lists"])
    html_ = re.sub(r"<(h[1-3])>", lambda m: f'<{m.group(1)} class="{ {"h1": "t1", "h2": "t2", "h3": "t3"}[m.group(1)]}">',
                   html_)
    return html_, fm


# ============================================================ texto puro
def ler_texto(caminho):
    dados = open(caminho, "rb").read()
    for cod in ("utf-8-sig", "utf-16"):
        if cod == "utf-16" and not dados[:2] in (b"\xff\xfe", b"\xfe\xff"):
            continue
        try:
            return dados.decode(cod)
        except UnicodeDecodeError:
            pass
    return dados.decode("cp1252", errors="replace")


def texto_para_html(t):
    return "".join(f"<p>{esc(l) or '<br>'}</p>" for l in t.replace("\r\n", "\n").split("\n"))


# ============================================================ RTF -> HTML
_PULAR = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "footer", "headerl", "headerr",
          "footerl", "footerr", "object", "fldinst", "themedata", "colorschememapping", "latentstyles",
          "datastore", "xmlnstbl", "listtable", "listoverridetable", "rsidtbl", "generator", "mmathPr",
          "pgdsctbl", "filetbl", "revtbl", "bkmkstart", "bkmkend", "nonshppict", "shppict", "wgrffmtfilter"}


def rtf_para_html(rtf):
    """RTF -> HTML (parágrafos com spans formatados). Cobre o que editores de texto produzem no dia a dia."""
    cores = [None]
    fontes = {}
    pilha = []
    st = {"b": False, "i": False, "u": False, "s": False, "fs": 24, "cf": 0, "hl": 0, "f": None,
          "pular": False, "uc": 1}
    par = {"align": "left", "li": 0, "sl": 0}
    paragrafos, atual = [], []
    i, n = 0, len(rtf)
    pular_chars = 0
    destino_fonte = None

    def emitir(texto):
        nonlocal pular_chars
        if st["pular"] or not texto:
            return
        if destino_fonte is not None:
            return
        estilo = []
        if st["fs"] != 24:
            estilo.append(f"font-size:{st['fs'] / 2:g}pt")
        if st["cf"] and st["cf"] < len(cores) and cores[st["cf"]]:
            estilo.append(f"color:{cores[st['cf']]}")
        if st["hl"] and st["hl"] < len(cores) and cores[st["hl"]]:
            estilo.append(f"background-color:{cores[st['hl']]}")
        if st["f"] is not None and fontes.get(st["f"]):
            estilo.append(f"font-family:'{fontes[st['f']]}'")
        deco = [d for d, k in (("underline", "u"), ("line-through", "s")) if st[k]]
        if deco:
            estilo.append("text-decoration:" + " ".join(deco))
        t = esc(texto).replace("\t", "&#9;")
        if st["b"]:
            t = f"<b>{t}</b>"
        if st["i"]:
            t = f"<i>{t}</i>"
        atual.append(f'<span style="{";".join(estilo)}">{t}</span>' if estilo else t)

    def fechar_paragrafo():
        estilo = []
        if par["align"] != "left":
            estilo.append(f"text-align:{par['align']}")
        if par["li"]:
            estilo.append(f"margin-left:{par['li'] / 15:.0f}px")        # twips -> px (1440/in, 96px/in)
        if par["sl"] and par["sl"] > 0:
            estilo.append(f"line-height:{par['sl'] / 240:.2f}")
        conteudo = "".join(atual) or "<br>"
        paragrafos.append(f'<p style="{";".join(estilo)}">{conteudo}</p>' if estilo else f"<p>{conteudo}</p>")
        atual.clear()

    buf = []

    def solta_buf():
        if buf:
            emitir("".join(buf))
            buf.clear()

    # tabela de fontes/cores: precisamos ler mesmo dentro dos grupos "pulados"
    while i < n:
        c = rtf[i]
        if c == "{":
            solta_buf()
            pilha.append(dict(st))
            i += 1
            continue
        if c == "}":
            solta_buf()
            if pilha:
                anterior = pilha.pop()
                if st.get("_fontes") and not anterior.get("_fontes"):
                    destino_fonte = None
                st.clear()
                st.update(anterior)
            i += 1
            continue
        if c == "\\":
            i += 1
            if i >= n:
                break
            c2 = rtf[i]
            if c2 in "\\{}":
                if pular_chars:
                    pular_chars -= 1
                else:
                    buf.append(c2)
                i += 1
                continue
            if c2 == "'":
                hexa = rtf[i + 1:i + 3]
                i += 3
                if pular_chars:
                    pular_chars -= 1
                    continue
                try:
                    buf.append(bytes([int(hexa, 16)]).decode("cp1252", errors="replace"))
                except ValueError:
                    pass
                continue
            if c2 == "*":
                st["pular"] = True
                i += 1
                continue
            if c2 == "~":
                buf.append("\u00a0")
                i += 1
                continue
            if c2 in "\n\r":
                solta_buf()
                fechar_paragrafo()
                i += 1
                continue
            m = re.match(r"([a-zA-Z]+)(-?\d+)? ?", rtf[i:])
            if not m:
                i += 1
                continue
            palavra, num = m.group(1), m.group(2)
            v = int(num) if num is not None else None
            i += m.end()
            if palavra in _PULAR:
                solta_buf()
                if palavra == "fonttbl":
                    st["_fontes"] = True
                    destino_fonte = ""
                elif palavra == "colortbl":
                    # lê a tabela de cores inteira aqui
                    fim = rtf.find("}", i)
                    corpo = rtf[i:fim]
                    cores[:] = []
                    for entrada in corpo.split(";")[:-1] if ";" in corpo else []:
                        r = re.search(r"\\red(\d+)", entrada)
                        g = re.search(r"\\green(\d+)", entrada)
                        b = re.search(r"\\blue(\d+)", entrada)
                        cores.append(f"#{int(r.group(1)):02x}{int(g.group(1)):02x}{int(b.group(1)):02x}" if r and g and b else None)
                    if not cores:
                        cores.append(None)
                    i = fim
                    st["pular"] = True
                else:
                    st["pular"] = True
                continue
            if st.get("_fontes"):
                if palavra == "f" and v is not None:
                    st["_f_atual"] = v
                    fim = rtf.find(";", i)
                    nome = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", rtf[i:fim]).strip()
                    if nome:
                        fontes[v] = nome
                continue
            if st["pular"]:
                continue
            textuais = {"u", "uc", "tab", "emdash", "endash", "bullet", "lquote", "rquote", "ldblquote", "rdblquote"}
            if palavra not in textuais:
                solta_buf()                  # só muda de trecho quando a FORMATAÇÃO muda
            if palavra == "par":
                fechar_paragrafo()
            elif palavra == "line":
                atual.append("<br>")
            elif palavra == "tab":
                buf.append("\t")
            elif palavra == "pard":
                par.update(align="left", li=0, sl=0)
            elif palavra == "plain":
                st.update(b=False, i=False, u=False, s=False, fs=24, cf=0, hl=0, f=None)
            elif palavra == "b":
                st["b"] = v != 0
            elif palavra == "i":
                st["i"] = v != 0
            elif palavra in ("ul", "uld", "uldb", "ulw"):
                st["u"] = v != 0
            elif palavra == "ulnone":
                st["u"] = False
            elif palavra in ("strike", "striked"):
                st["s"] = v != 0
            elif palavra == "fs" and v:
                st["fs"] = v
            elif palavra == "cf":
                st["cf"] = v or 0
            elif palavra in ("highlight", "cb", "chcbpat"):
                st["hl"] = v or 0
            elif palavra == "f":
                st["f"] = v
            elif palavra == "qc":
                par["align"] = "center"
            elif palavra == "qr":
                par["align"] = "right"
            elif palavra == "qj":
                par["align"] = "justify"
            elif palavra == "ql":
                par["align"] = "left"
            elif palavra == "li":
                par["li"] = max(0, v or 0)
            elif palavra == "sl" and v:
                par["sl"] = abs(v)
            elif palavra == "uc":
                st["uc"] = v or 0
            elif palavra == "u" and v is not None:
                buf.append(chr(v + 65536 if v < 0 else v))
                pular_chars = st["uc"]
            elif palavra in ("emdash",):
                buf.append("—")
            elif palavra in ("endash",):
                buf.append("–")
            elif palavra in ("bullet",):
                buf.append("•")
            elif palavra in ("lquote", "rquote"):
                buf.append("‘" if palavra == "lquote" else "’")
            elif palavra in ("ldblquote", "rdblquote"):
                buf.append("“" if palavra == "ldblquote" else "”")
            continue
        if c in "\r\n":
            i += 1
            continue
        if pular_chars:
            pular_chars -= 1
        elif st.get("_fontes") or st["pular"]:
            pass
        else:
            buf.append(c)
        i += 1
    solta_buf()
    if atual:
        fechar_paragrafo()
    return "".join(paragrafos)


# ============================================================ HTML -> blocos (pra RTF e DOCX)
class _Blocos(HTMLParser):
    """Lê o HTML do editor e produz blocos: {tipo, align, indent_px, line_height, runs:[(texto, fmt)]}.
    tipo: p, h1, h2, h3, li-ul, li-ol, img."""
    BLOCOS = {"p", "div", "h1", "h2", "h3", "h4", "li", "blockquote", "pre"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocos = []
        self.pilha_fmt = [{}]
        self.listas = []
        self.bloco = None
        self.tags = []

    def _novo(self, tipo, estilo):
        self._fecha()
        b = {"tipo": tipo, "runs": [], "align": None, "indent": 0, "lh": None, "classe": ""}
        b.update(estilo)
        self.bloco = b

    def _fecha(self):
        if self.bloco is not None:
            self.blocos.append(self.bloco)
            self.bloco = None

    @staticmethod
    def _css(s):
        d = {}
        for parte in (s or "").split(";"):
            if ":" in parte:
                k, v = parte.split(":", 1)
                d[k.strip().lower()] = v.strip()
        return d

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        css = self._css(a.get("style"))
        self.tags.append(tag)
        if tag in ("ul", "ol"):
            self.listas.append(tag)
            return
        if tag in self.BLOCOS:
            tipo = tag
            if tag == "li":
                tipo = "li-" + (self.listas[-1] if self.listas else "ul")
            elif tag in ("div", "blockquote", "pre", "h4"):
                tipo = "p" if tag != "h4" else "h3"
            est = {"align": css.get("text-align"), "classe": a.get("class") or ""}
            ml = css.get("margin-left") or css.get("padding-left")
            if ml:
                m = re.match(r"([\d.]+)(px|pt|em)?", ml)
                if m:
                    vv = float(m.group(1))
                    est["indent"] = vv * (1.333 if m.group(2) == "pt" else 16 if m.group(2) == "em" else 1)
            if css.get("line-height"):
                try:
                    est["lh"] = float(css["line-height"])
                except ValueError:
                    pass
            self._novo(tipo, est)
            self.pilha_fmt.append(dict(self.pilha_fmt[-1]))
            return
        if tag == "br":
            if self.bloco is None:
                self._novo("p", {})
            self.bloco["runs"].append(("\n", dict(self.pilha_fmt[-1])))
            self.tags.pop()
            return
        if tag == "img":
            src = a.get("src") or ""
            self._fecha()
            self.blocos.append({"tipo": "img", "src": src, "runs": [], "align": None, "indent": 0, "lh": None,
                                "classe": ""})
            self.tags.pop()
            return
        f = dict(self.pilha_fmt[-1])
        if tag in ("b", "strong"):
            f["b"] = True
        elif tag in ("i", "em"):
            f["i"] = True
        elif tag == "u":
            f["u"] = True
        elif tag in ("s", "strike", "del"):
            f["s"] = True
        elif tag in ("code", "kbd"):
            f["font"] = "monospace"
        elif tag == "font":
            if a.get("color"):
                f["color"] = a["color"]
            if a.get("face"):
                f["font"] = a["face"]
        if css.get("font-weight") in ("bold", "700", "800", "900"):
            f["b"] = True
        if css.get("font-style") == "italic":
            f["i"] = True
        deco = css.get("text-decoration") or css.get("text-decoration-line") or ""
        if "underline" in deco:
            f["u"] = True
        if "line-through" in deco:
            f["s"] = True
        if css.get("color"):
            f["color"] = css["color"]
        if css.get("background-color"):
            f["bg"] = css["background-color"]
        if css.get("font-size"):
            m = re.match(r"([\d.]+)(pt|px)?", css["font-size"])
            if m:
                f["size"] = float(m.group(1)) * (0.75 if m.group(2) == "px" else 1)
        if css.get("font-family"):
            f["font"] = css["font-family"].split(",")[0].strip("'\" ")
        self.pilha_fmt.append(f)

    def handle_endtag(self, tag):
        if self.tags and self.tags[-1] == tag:
            self.tags.pop()
        if tag in ("ul", "ol"):
            if self.listas:
                self.listas.pop()
            return
        if tag in self.BLOCOS:
            self._fecha()
        if len(self.pilha_fmt) > 1 and tag not in ("br", "img"):
            self.pilha_fmt.pop()

    def handle_data(self, data):
        if not data:
            return
        if self.bloco is None:
            if not data.strip():
                return
            self._novo("p", {})
        self.bloco["runs"].append((data.replace("\n", " ") if "pre" not in self.tags else data,
                                   dict(self.pilha_fmt[-1])))


def html_blocos(html_):
    p = _Blocos()
    p.feed(html_ or "")
    p.close()
    p._fecha()
    return p.blocos


def _cor_hex(c):
    c = (c or "").strip()
    m = re.match(r"#([0-9a-fA-F]{6})$", c)
    if m:
        return m.group(1).lower()
    m = re.match(r"#([0-9a-fA-F]{3})$", c)
    if m:
        return "".join(x * 2 for x in m.group(1)).lower()
    m = re.match(r"rgba?\((\d+),\s*(\d+),\s*(\d+)", c)
    if m:
        return "".join(f"{int(x):02x}" for x in m.groups())
    return None


TAMANHO_BLOCO = {"h1": 20, "h2": 16, "h3": 14}
CLASSE_TAMANHO = {"titulo": 26, "subtitulo": 16, "t1": 20, "t2": 16, "t3": 14}


# ============================================================ HTML -> RTF
def html_para_rtf(html_, fonte_padrao="Adwaita Sans", tamanho_padrao=11):
    blocos = html_blocos(html_)
    fontes = [fonte_padrao]
    cores = []

    def fid(nome):
        if nome not in fontes:
            fontes.append(nome)
        return fontes.index(nome)

    def cid(c):
        h = _cor_hex(c)
        if not h:
            return None
        if h not in cores:
            cores.append(h)
        return cores.index(h) + 1

    def rtf_texto(t):
        out = []
        for ch in t:
            o = ord(ch)
            if ch in "\\{}":
                out.append("\\" + ch)
            elif ch == "\t":
                out.append("\\tab ")
            elif ch == "\n":
                out.append("\\line ")
            elif o < 128:
                out.append(ch)
            else:
                out.append(f"\\u{o if o < 32768 else o - 65536}?")
        return "".join(out)

    corpo = []
    num = 0
    for b in blocos:
        if b["tipo"] == "img":
            continue
        pard = ["\\pard"]
        al = b.get("align")
        pard.append({"center": "\\qc", "right": "\\qr", "justify": "\\qj"}.get(al, "\\ql"))
        indent = b.get("indent") or 0
        lista = b["tipo"].startswith("li-")
        if lista:
            indent += 24
        if indent:
            pard.append(f"\\li{int(indent * 15)}")
        if b.get("lh"):
            pard.append(f"\\sl{int(b['lh'] * 240)}\\slmult1")
        tam_bloco = TAMANHO_BLOCO.get(b["tipo"])
        for cls, t in CLASSE_TAMANHO.items():
            if cls in (b.get("classe") or "").split():
                tam_bloco = t
        prefixo = ""
        if b["tipo"] == "li-ul":
            prefixo = "\\bullet\\tab "
            num = 0
        elif b["tipo"] == "li-ol":
            num += 1
            prefixo = f"{num}.\\tab "
        else:
            num = 0
        partes = []
        for texto, f in b["runs"]:
            fmt = [f"\\f{fid(f['font'])}" if f.get("font") else "\\f0"]
            tam = f.get("size") or tam_bloco or tamanho_padrao
            fmt.append(f"\\fs{int(round(tam * 2))}")
            fmt.append("\\b" if f.get("b") or b["tipo"] in ("h1", "h2", "h3") else "\\b0")
            fmt.append("\\i" if f.get("i") else "\\i0")
            fmt.append("\\ul" if f.get("u") else "\\ulnone")
            fmt.append("\\strike" if f.get("s") else "\\strike0")
            c = cid(f.get("color")) if f.get("color") else None
            fmt.append(f"\\cf{c}" if c else "\\cf0")
            h = cid(f.get("bg")) if f.get("bg") else None
            fmt.append(f"\\highlight{h}" if h else "\\highlight0")
            partes.append("{" + "".join(fmt) + " " + rtf_texto(texto) + "}")
        corpo.append("".join(pard) + " " + prefixo + "".join(partes) + "\\par")
    tabela_fontes = "".join(f"{{\\f{i}\\fnil {n};}}" for i, n in enumerate(fontes))
    tabela_cores = ";" + "".join(f"\\red{int(h[0:2], 16)}\\green{int(h[2:4], 16)}\\blue{int(h[4:6], 16)};"
                                 for h in cores)
    return ("{\\rtf1\\ansi\\ansicpg1252\\deff0\\uc1{\\fonttbl" + tabela_fontes + "}{\\colortbl" + tabela_cores +
            "}\\viewkind4\n" + "\n".join(corpo) + "\n}")


# ============================================================ DOCX
def docx_para_html(caminho):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = docx.Document(caminho)
    al = {WD_ALIGN_PARAGRAPH.CENTER: "center", WD_ALIGN_PARAGRAPH.RIGHT: "right",
          WD_ALIGN_PARAGRAPH.JUSTIFY: "justify"}
    saida, lista_atual = [], None

    def runs(p):
        out = []
        for r in p.runs:
            t = esc(r.text).replace("\t", "&#9;").replace("\n", "<br>")
            if not t:
                continue
            est = []
            if r.font.size:
                est.append(f"font-size:{r.font.size.pt:g}pt")
            if r.font.color is not None and r.font.color.type is not None and r.font.color.rgb is not None:
                est.append(f"color:#{r.font.color.rgb}")
            if r.font.highlight_color:
                cor = {3: "#00ff00", 4: "#ff00ff", 6: "#ff0000", 7: "#ffff00", 5: "#0000ff", 2: "#000000"}.get(
                    int(r.font.highlight_color), "#ffff00")
                est.append(f"background-color:{cor}")
            if r.font.name:
                est.append(f"font-family:'{r.font.name}'")
            deco = [x for x, on in (("underline", r.underline), ("line-through", r.font.strike)) if on]
            if deco:
                est.append("text-decoration:" + " ".join(deco))
            if r.bold:
                t = f"<b>{t}</b>"
            if r.italic:
                t = f"<i>{t}</i>"
            out.append(f'<span style="{";".join(est)}">{t}</span>' if est else t)
        return "".join(out) or "<br>"

    def estilo_par(p):
        est = []
        if p.alignment in al:
            est.append(f"text-align:{al[p.alignment]}")
        pf = p.paragraph_format
        if pf.left_indent:
            est.append(f"margin-left:{pf.left_indent.pt * 1.333:.0f}px")
        if pf.line_spacing and isinstance(pf.line_spacing, float):
            est.append(f"line-height:{pf.line_spacing:g}")
        return f' style="{";".join(est)}"' if est else ""

    for p in d.paragraphs:
        nome = (p.style.name or "").lower() if p.style is not None else ""
        if "list bullet" in nome or "list number" in nome or "list paragraph" in nome:
            tag = "ol" if "number" in nome else "ul"
            if lista_atual != tag:
                if lista_atual:
                    saida.append(f"</{lista_atual}>")
                saida.append(f"<{tag}>")
                lista_atual = tag
            saida.append(f"<li>{runs(p)}</li>")
            continue
        if lista_atual:
            saida.append(f"</{lista_atual}>")
            lista_atual = None
        if nome == "title":
            saida.append(f'<h1 class="titulo"{estilo_par(p)}>{runs(p)}</h1>')
        elif nome == "subtitle":
            saida.append(f'<h2 class="subtitulo"{estilo_par(p)}>{runs(p)}</h2>')
        elif nome.startswith("heading"):
            nivel = min(3, int(re.sub(r"\D", "", nome) or 1))
            saida.append(f'<h{nivel} class="t{nivel}"{estilo_par(p)}>{runs(p)}</h{nivel}>')
        else:
            saida.append(f"<p{estilo_par(p)}>{runs(p)}</p>")
    if lista_atual:
        saida.append(f"</{lista_atual}>")
    for t in d.tables:                       # tabelas: o editor não tem tabela -> viram linhas de texto
        for linha in t.rows:
            saida.append("<p>" + " &#9; ".join(esc(c.text) for c in linha.cells) + "</p>")
    return "".join(saida)


def html_para_docx(html_, caminho, fonte_padrao="Adwaita Sans", tamanho_padrao=11):
    import io
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Mm
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    normal = d.styles["Normal"]
    normal.font.name = fonte_padrao
    normal.font.size = Pt(tamanho_padrao)
    al = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
          "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    for b in html_blocos(html_):
        if b["tipo"] == "img":
            m = re.match(r"data:image/[\w+]+;base64,(.*)", b.get("src") or "", re.S)
            if m:
                try:
                    d.add_picture(io.BytesIO(base64.b64decode(m.group(1))), width=Mm(150))
                except Exception:  # noqa: BLE001
                    pass
            continue
        classes = (b.get("classe") or "").split()
        if "titulo" in classes:
            p = d.add_paragraph(style="Title")
        elif "subtitulo" in classes:
            p = d.add_paragraph(style="Subtitle")
        elif b["tipo"] in ("h1", "h2", "h3"):
            p = d.add_heading(level=int(b["tipo"][1]))
        elif b["tipo"] == "li-ul":
            p = d.add_paragraph(style="List Bullet")
        elif b["tipo"] == "li-ol":
            p = d.add_paragraph(style="List Number")
        else:
            p = d.add_paragraph()
        if b.get("align") in al:
            p.alignment = al[b["align"]]
        if b.get("indent"):
            p.paragraph_format.left_indent = Pt(b["indent"] * 0.75)
        if b.get("lh"):
            p.paragraph_format.line_spacing = b["lh"]
        for texto, f in b["runs"]:
            r = p.add_run(texto.replace("\n", "\v") if "\n" in texto else texto)
            if "\n" in texto:
                r.text = ""
                partes = texto.split("\n")
                for k, parte in enumerate(partes):
                    r.add_text(parte)
                    if k < len(partes) - 1:
                        r.add_break()
            r.bold = bool(f.get("b")) or None
            r.italic = bool(f.get("i")) or None
            r.underline = bool(f.get("u")) or None
            r.font.strike = bool(f.get("s")) or None
            if f.get("size"):
                r.font.size = Pt(f["size"])
            if f.get("font") and f["font"] != "monospace":
                r.font.name = f["font"]
            elif f.get("font") == "monospace":
                r.font.name = "Adwaita Mono"
            h = _cor_hex(f.get("color"))
            if h:
                r.font.color.rgb = RGBColor.from_string(h.upper())
            if _cor_hex(f.get("bg")):
                from docx.enum.text import WD_COLOR_INDEX
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    d.save(caminho)


def docx_texto(caminho):
    import docx
    return "\n".join(p.text for p in docx.Document(caminho).paragraphs)


# ============================================================ prévia (tela inicial)
def preview(caminho, linhas=40):
    ext = os.path.splitext(caminho)[1].lower()
    try:
        if ext == ".docx":
            t = docx_texto(caminho)
        elif ext == ".rtf":
            t = re.sub(r"<[^>]+>", "", rtf_para_html(ler_texto(caminho)).replace("</p>", "\n"))
            t = html.unescape(t)
        else:
            t = ler_texto(caminho)
            if ext == ".artry" and re.match(r"\s*<", t):
                t = re.sub(r"<(br|/p|/h\d|/li|/div)[^>]*>", "\n", t, flags=re.I)
                t = html.unescape(re.sub(r"<[^>]+>", "", t))
            elif ext in (".hyp", ".md"):
                _, t, _ = hyp_partes(t)
    except Exception:  # noqa: BLE001
        t = ""
    return t.replace("\r\n", "\n").split("\n")[:linhas]
